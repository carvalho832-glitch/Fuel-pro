"""Fuel Pro: importação experimental de preços por posto da ANP.

Modo de uso:
    python import_anp.py --arquivo pesquisa.csv
    python import_anp.py --arquivo pesquisa.xlsx --geocode --max-geocode 15
    python import_anp.py --mais-recente --geocode --max-geocode 15

Importante:
- Apenas observações medidas, nunca preços calculados.
- Geocodificação do Nominatim é aproximada, limitada e armazenada em cache.
- Sem postos corretamente identificados/localizados, NÃO sobrescreve dados.json.
"""
from __future__ import annotations

import argparse
import csv
from datetime import datetime
from decimal import Decimal, InvalidOperation
import io
import json
from pathlib import Path
import re
import time
import unicodedata
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup

ANP_WEEKLY_PAGE = (
    "https://www.gov.br/anp/pt-br/assuntos/precos-e-defesa-da-concorrencia/"
    "precos/levantamento-de-precos-de-combustiveis-ultimas-semanas-pesquisadas"
)
UA = "FuelPro-Research/0.2 (https://github.com/carvalho832-glitch/Fuel-pro)"
CITY_BOUNDS = {
    "SAO JOSE DOS CAMPOS": (-23.40, -22.95, -46.20, -45.60),
    "ARARAQUARA": (-21.95, -21.60, -48.45, -47.95),
}
FUELS = {
    "GASOLINA": "gas",
    "GASOLINA COMUM": "gas",
    "GASOLINA ADITIVADA": "gas_ad",
    "ETANOL": "eta",
    "ETANOL HIDRATADO": "eta",
    "ETANOL ADITIVADO": "eta_ad",
    "DIESEL": "die",
    "OLEO DIESEL": "die",
    "DIESEL S500": "die",
    "DIESEL S10": "die_s10",
    "GNV": "gnv",
}
FIELDS = {
    "municipio": ["MUNICIPIO", "CIDADE"],
    "uf": ["ESTADO - SIGLA", "UF", "ESTADO"],
    "cnpj": ["CNPJ DA REVENDA", "CNPJ"],
    "name": ["REVENDA", "RAZAO SOCIAL", "RAZAO"],
    "product": ["PRODUTO", "COMBUSTIVEL"],
    "price": ["VALOR DE VENDA", "PRECO DE VENDA", "PRECO DE REVENDA", "PRECO"],
    "date": ["DATA DA COLETA", "DATA COLETA", "DATA DE COLETA"],
    "street": ["NOME DA RUA", "LOGRADOURO", "ENDERECO"],
    "number": ["NUMERO RUA", "NUMERO", "NUMERO DO LOGRADOURO"],
    "district": ["BAIRRO"],
    "zip": ["CEP"],
}

def norm(value):
    value = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"\s+", " ", "".join(c for c in value if not unicodedata.combining(c)).upper()).strip()

def key_header(value):
    return re.sub(r"[^A-Z0-9]+", " ", norm(value)).strip()

def pick(row, field):
    normalized = {key_header(k): v for k, v in row.items() if k}
    for name in FIELDS[field]:
        key = key_header(name)
        if key in normalized:
            return normalized[key]
    return ""

def valid_header(row):
    return (bool(pick(row, "cnpj")) and bool(pick(row, "product"))
            and bool(pick(row, "price")))

def csv_rows(binary):
    # ANP exports commonly use semicolon-delimited latin-1 or UTF-8.
    try:
        txt = binary.decode("utf-8-sig")
    except UnicodeDecodeError:
        txt = binary.decode("cp1252")
    sample = txt[:2000]
    sep = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.DictReader(io.StringIO(txt), delimiter=sep)
    for row in reader:
        yield row

def xlsx_rows(binary):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(binary), read_only=True, data_only=True)
    for sheet in wb.worksheets:
        header = None
        for values in sheet.iter_rows(values_only=True):
            if not values:
                continue
            if header is None:
                names = [str(v or "") for v in values]
                if any(key_header(n) in ("CNPJ DA REVENDA", "CNPJ") for n in names):
                    header = names
                continue
            yield dict(zip(header, values))

def parse_rows(rows, cities=CITY_BOUNDS):
    grouped = {}
    for row in rows:
        city = norm(pick(row, "municipio"))
        if city not in cities or norm(pick(row, "uf")) != "SP":
            continue
        cnpj = re.sub(r"\D", "", str(pick(row, "cnpj")))
        name = str(pick(row, "name") or "").strip()
        fuel = FUELS.get(norm(pick(row, "product")))
        if len(cnpj) != 14 or len(name) < 4 or fuel is None:
            continue
        try:
            value = Decimal(str(pick(row, "price")).replace(",", ".").strip())
        except (InvalidOperation, ValueError):
            continue
        if not (Decimal("0.50") < value < Decimal("30.00")):
            continue
        date_raw = pick(row, "date")
        if isinstance(date_raw, datetime):
            date_iso = date_raw.date().isoformat()
        else:
            try:
                date_iso = datetime.strptime(str(date_raw).strip()[:10], "%d/%m/%Y").date().isoformat()
            except ValueError:
                try:
                    date_iso = datetime.strptime(str(date_raw).strip()[:10], "%Y-%m-%d").date().isoformat()
                except ValueError:
                    continue  # Sem data comprovável, não existe preço exibível.
        street = str(pick(row, "street") or "").strip()
        number = str(pick(row, "number") or "").strip()
        district = str(pick(row, "district") or "").strip()
        address = ", ".join(s for s in (street, number, district, city.title() + " - SP") if s)
        if len(street) < 4:
            continue
        p = grouped.setdefault(cnpj, {
            "id": cnpj, "name": name, "address": address, "city": city,
            "source": "ANP - levantamento semanal", "cnpj": cnpj,
            "prices": {}, "collectedAtByFuel": {}, "locationPrecision": "approximate"
        })
        if (fuel not in p["collectedAtByFuel"] or date_iso > p["collectedAtByFuel"][fuel]):
            p["prices"][fuel] = float(value)
            p["collectedAtByFuel"][fuel] = date_iso
    return list(grouped.values())

def latest_anp_file(session):
    r = session.get(ANP_WEEKLY_PAGE, timeout=30)
    r.raise_for_status()
    soup = BeautifulSoup(r.text, "html.parser")
    for anchor in soup.find_all("a", href=True):
        if "PRECOS POR POSTO REVENDEDOR" in norm(anchor.get_text(" ", strip=True)):
            url = urljoin(r.url, anchor["href"])
            if not url.startswith(("https://www.gov.br/", "https://www.anp.gov.br/")):
                raise RuntimeError("URL de download não pertence ao site oficial ANP: " + url)
            file_response = session.get(url, timeout=90)
            file_response.raise_for_status()
            print("ANP arquivo:", file_response.url, "| content-type:", file_response.headers.get("content-type"),
                  "| bytes:", len(file_response.content), "| assinatura:", repr(file_response.content[:20]))
            return file_response.content
    raise RuntimeError("A planilha de preços por posto não foi encontrada na página oficial da ANP")

def has_city_coordinate(city, lat, lng):
    lo_lat, hi_lat, lo_lng, hi_lng = CITY_BOUNDS[city]
    return lo_lat <= lat <= hi_lat and lo_lng <= lng <= hi_lng

def geocode(station, cache, session):
    cnpj = station["cnpj"]
    if cnpj in cache:
        item = cache[cnpj]
        if item and has_city_coordinate(station["city"], item["lat"], item["lng"]):
            return item
        return None
    query = station["address"].replace(" - SP", ", São Paulo, Brasil")
    response = session.get(
        "https://nominatim.openstreetmap.org/search",
        params={"q": query, "format": "json", "limit": 1, "countrycodes": "br"},
        timeout=25, headers={"User-Agent": UA}
    )
    response.raise_for_status()
    matches = response.json()
    cache[cnpj] = None  # Evita repetir buscas que não retornaram resultado.
    if matches:
        lat, lng = float(matches[0]["lat"]), float(matches[0]["lon"])
        if has_city_coordinate(station["city"], lat, lng):
            cache[cnpj] = {"lat": lat, "lng": lng}
    return cache[cnpj]

def run(args, session=None):
    session = session or requests.Session()
    session.headers.update({"User-Agent": UA})
    if args.arquivo:
        binary = Path(args.arquivo).read_bytes()
    else:
        binary = latest_anp_file(session)
    rows = xlsx_rows(binary) if binary[:2] == b"PK" else csv_rows(binary)
    if getattr(args, "apenas_validar", False):
        from itertools import chain, islice
        sample = list(islice(rows, 3))
        print("ANP amostra de colunas:", [list(s.keys()) for s in sample[:1]])
        print("ANP amostra cidade/UF:", [(pick(s, "municipio"), pick(s, "uf"), pick(s, "product")) for s in sample])
        rows = chain(sample, rows)
    stations = parse_rows(rows)
    if not stations:
        raise RuntimeError("Importação interrompida: nenhum posto válido de SJC/Araraquara na fonte.")
    if getattr(args, "apenas_validar", False):
        print(f"Fonte ANP validada: {len(stations)} postos identificados nas cidades piloto; nenhum dado publicado.")
        return len(stations)
    cache_path = Path(args.cache)
    cache = json.loads(cache_path.read_text(encoding="utf-8")) if cache_path.exists() else {}
    mapped = []
    requests_count = 0
    for p in stations:
        if p["cnpj"] in cache:
            coords = cache[p["cnpj"]]
        elif args.geocode and requests_count < args.max_geocode:
            time.sleep(1.2)  # Respeita o limite público do Nominatim; guarda tudo em cache.
            coords = geocode(p, cache, session)
            requests_count += 1
        else:
            coords = None
        if coords and has_city_coordinate(p["city"], coords["lat"], coords["lng"]):
            p["lat"], p["lng"] = coords["lat"], coords["lng"]
            mapped.append(p)
    cache_path.write_text(json.dumps(cache, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"ANP: {len(stations)} postos válidos; {len(mapped)} com localização; {requests_count} geocodificados.")
    if not mapped:
        raise RuntimeError("Sem postos mapeáveis. Arquivo de produção preservado, sem coordenadas inventadas.")
    output = Path(args.output)
    temp = output.with_suffix(".tmp")
    temp.write_text(json.dumps(mapped, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(output)
    print("Importação concluída. Observações reais e datadas, localização aproximada:", output)

def main():
    parser = argparse.ArgumentParser()
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--arquivo", help="CSV ou XLSX baixado da ANP")
    group.add_argument("--mais-recente", action="store_true", help="Busca a última planilha semanal na ANP")
    parser.add_argument("--geocode", action="store_true", help="Geocodificação aproximada com cache")
    parser.add_argument("--max-geocode", type=int, default=15, help="Limite de novas buscas por execução")
    parser.add_argument("--cache", default="geocache.json")
    parser.add_argument("--output", default="dados.json")
    parser.add_argument("--apenas-validar", action="store_true", help="Valida planilha ANP sem alterar dados")
    run(parser.parse_args())

if __name__ == "__main__":
    main()
